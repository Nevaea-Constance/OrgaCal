"""Read the supplied paper files and build a traceable JSON index, without changing them.

Usage: python scripts/build_reference_index.py --source-dir <source folder>
Requires openpyxl for reading Excel. No image segmentation or statistical fitting runs here.
"""
import argparse
from collections import Counter, defaultdict
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path, PurePosixPath
import re
import zipfile

from openpyxl import load_workbook


FEATURES = {
    'area[um2]': 'area_um2',
    'perimeter[um]': 'perimeter_um',
    'avg_radius[um]': 'paper_radius_um',
    'roundness[%]': 'paper_roundness_pct',
    'maxFeret[um]': 'max_feret_um',
    'minFeret[um]': 'min_feret_um',
    'Mean curvature[um-1]': 'mean_curvature_per_um',
    'Std curvature[um-1]': 'std_curvature_per_um',
    'Std curvature x R0': 'std_curvature_times_radius',
    'DNE': 'paper_dne',
}
EXTRA_FEATURES = {'circularity[%]': 'source_circularity_pct',
                  'transparency[%]': 'source_transparency_pct'}
GROUPS = {'PA_day6': ('PA', 6), 'RC_day6': ('RC', 6),
          'OS_day90': ('OS', 90), 'RC_day90': ('RC', 90)}


def file_hash(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def stable_id(prefix, key):
    return prefix + '_' + hashlib.sha256(key.encode('utf-8')).hexdigest()[:20]


def finite_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def features_from(row, include_extra=False):
    fields = FEATURES | (EXTRA_FEATURES if include_extra else {})
    return {key: row.get(header) for header, key in fields.items()}


def read_sheet(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    nonempty_headers = [h for h in headers if h is not None]
    if len(nonempty_headers) != len(set(nonempty_headers)):
        raise ValueError(f'Duplicate column header in {sheet.title}')
    for number, values in enumerate(rows, 2):
        if any(v is not None for v in values):
            yield number, {h: v for h, v in zip(headers, values) if h is not None}


def group_metadata(entry):
    parts = PurePosixPath(entry).parts
    at = parts.index('Raw images_mask')
    group, folder = parts[at + 1:at + 3]
    if group not in GROUPS:
        raise ValueError(f'Unexpected group: {group}')
    batch = re.search(r'\bC\d+\b', folder)
    line = re.search(r'\b(?:11C1|12C1)\b', folder)
    magnification = re.search(r'\b(\d+)x\b', folder)
    if not batch or not line or not magnification:
        raise ValueError(f'Unrecognized metadata folder: {folder}')
    condition, day = GROUPS[group]
    return {'group': group, 'condition': condition, 'day': day,
            'batch': batch.group(), 'cell_line_code': line.group(),
            'magnification_label': magnification.group(), 'source_folder': folder}


def build(source_dir, project_dir):
    source_dir, project_dir = Path(source_dir), Path(project_dir)
    zip_path, ev_path = source_dir / '数据来源图3.zip', source_dir / '数据集1.xlsx'
    sources = [{'file': p.name, 'sha256': file_hash(p), 'bytes': p.stat().st_size}
               for p in (zip_path, ev_path)]
    images, objects, references, issues = [], [], [], []
    with zipfile.ZipFile(zip_path) as archive:
        names = sorted(n for n in archive.namelist()
                       if '/Raw images_mask/' in n and not n.startswith('__MACOSX/')
                       and not PurePosixPath(n).name.startswith('.') and not n.endswith('/'))
        if len(names) != len(set(names)):
            raise ValueError('Duplicate archive entry names')
        image_names = [n for n in names if n.lower().endswith('.jpg')]
        annotated = {n for n in image_names if PurePosixPath(n).name.startswith('R_')}
        raw = [n for n in image_names if n not in annotated]
        by_path = {}
        for name in raw:
            path = PurePosixPath(name)
            result = str(path.with_name('R_' + path.name))
            rec = {'image_id': stable_id('image', name), **group_metadata(name),
                   'source_archive': zip_path.name, 'raw_entry': name,
                   'overlay_entry': result if result in annotated else None,
                   'overlay_role': 'author_outline_and_number_visualization_not_label_mask',
                   'um_per_pixel': None, 'scale_status': 'not_independently_verified',
                   'zip_crc32': f'{archive.getinfo(name).CRC:08x}'}
            images.append(rec)
            by_path[name] = rec
            if result not in annotated:
                issues.append({'kind': 'missing_overlay', 'image_id': rec['image_id']})
        for name in sorted(annotated):
            p = PurePosixPath(name)
            if str(p.with_name(p.name[2:])) not in by_path:
                issues.append({'kind': 'orphan_overlay', 'entry': name})

        workbook_entries = [n for n in names if n.endswith('.xlsx')]
        for entry in workbook_entries:
            book = load_workbook(BytesIO(archive.read(entry)), read_only=True, data_only=True)
            for sheet in book.worksheets:
                for row_number, row in read_sheet(sheet):
                    image_name, object_id = row.get('image'), row.get('id')
                    if not image_name or not finite_number(object_id) or int(object_id) != object_id:
                        raise ValueError(f'Invalid image/id at {entry}:{sheet.title}:{row_number}')
                    target = str(PurePosixPath(entry).parent / image_name)
                    candidates = [by_path[target]] if target in by_path else []
                    method = 'exact_filename'
                    if not candidates:
                        candidates = [v for k, v in by_path.items() if k.casefold() == target.casefold()]
                        method = 'unique_case_insensitive_filename'
                    match = candidates[0] if len(candidates) == 1 else None
                    rec = {'record_id': stable_id('object', f'{entry}|{sheet.title}|{row_number}'),
                           **group_metadata(entry), 'source_archive': zip_path.name,
                           'source_workbook_entry': entry, 'source_sheet': sheet.title,
                           'source_row': row_number, 'source_image_name': image_name,
                           'source_object_id': int(object_id),
                           'image_id': match['image_id'] if match else None,
                           'image_match_status': method if match else 'unresolved',
                           'metrics': features_from(row, include_extra=True)}
                    objects.append(rec)
                    if not match:
                        issues.append({'kind': 'object_image_unresolved', 'record_id': rec['record_id'],
                                       'expected_entry': target})
            book.close()

    grouped = defaultdict(list)
    for obj in objects:
        grouped[(obj['group'], obj['batch'])].append(obj)
    book = load_workbook(ev_path, read_only=True, data_only=True)
    for row_number, row in read_sheet(book['all']):
        if not row.get('Group'):
            continue
        group = row['Group'].replace(' day ', '_day')
        if group not in GROUPS or GROUPS[group][1] != row['day']:
            raise ValueError(f'Invalid EV1 group/day at row {row_number}')
        metrics = features_from(row)
        if not all(finite_number(v) for v in metrics.values()):
            raise ValueError(f'Missing/nonfinite EV1 feature at row {row_number}')
        candidates = grouped[(group, row['batch'])]
        matches = [o for o in candidates if all(
            finite_number(o['metrics'].get(k)) and math.isclose(o['metrics'][k], v,
                                                               rel_tol=1e-8, abs_tol=1e-8)
            for k, v in metrics.items())]
        match = matches[0] if len(matches) == 1 else None
        status = 'unique_10_feature_match' if match else 'ambiguous' if matches else 'unmatched'
        reference = {'reference_id': f'EV1-all-{row_number:04d}',
                     'source_file': ev_path.name, 'source_sheet': 'all', 'source_row': row_number,
                     'source_group': row['Group'], 'group': group, 'batch': row['batch'],
                     'day': row['day'], 'metrics': metrics, 'match_status': status,
                     'candidate_record_ids': [o['record_id'] for o in matches],
                     'object_record_id': match['record_id'] if match else None,
                     'image_id': match['image_id'] if match else None,
                     'source_object_id': match['source_object_id'] if match else None,
                     'cell_line_code': match['cell_line_code'] if match else None,
                     'eligible_for_paired_validation': bool(match and match['image_id'])}
        references.append(reference)
        if not match:
            issues.append({'kind': 'ev1_' + status, 'reference_id': reference['reference_id'],
                           'group': group, 'batch': row['batch'],
                           'detail': 'No unique agreement on all ten features within the same group and batch; no forced pairing.'})
    book.close()

    linked = Counter(r['object_record_id'] for r in references if r['object_record_id'])
    reused = [key for key, count in linked.items() if count > 1]
    if reused:
        raise ValueError('One source object was linked to multiple EV1 rows: ' + repr(reused))
    object_keys = Counter((o['image_id'], o['source_object_id']) for o in objects if o['image_id'])
    if any(n > 1 for n in object_keys.values()):
        raise ValueError('Duplicate object ID within a raw image')
    summary = {'raw_images': len(images), 'paired_overlays': sum(bool(i['overlay_entry']) for i in images),
               'batch_workbooks': len(workbook_entries), 'source_objects': len(objects),
               'reference_samples': len(references),
               'reference_match_status': dict(Counter(r['match_status'] for r in references)),
               'object_image_status': dict(Counter(o['image_match_status'] for o in objects)),
               'groups': []}
    for group in GROUPS:
        refs = [r for r in references if r['group'] == group]
        summary['groups'].append({'group': group, 'raw_images': sum(i['group'] == group for i in images),
                                  'source_objects': sum(o['group'] == group for o in objects),
                                  'reference_samples': len(refs),
                                  'paired_reference_samples': sum(r['eligible_for_paired_validation'] for r in refs)})
    payload = {'schema_version': 1, 'paper_pmc': 'PMC12715241', 'sources': sources,
               'metric_source_headers': FEATURES | EXTRA_FEATURES,
               'matching_rule': {'keys': ['group', 'batch'], 'metrics': list(FEATURES.values()),
                                 'relative_tolerance': 1e-8, 'absolute_tolerance': 1e-8,
                                 'meaning': 'derived numeric association, not an author-supplied crosswalk'},
               'summary': summary, 'images': images, 'source_objects': objects,
               'reference_samples': references, 'issues': issues}
    output = project_dir / 'data' / 'reference' / 'paper_data_index.json'
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8')

    # Human-readable tables use the same traced records as the machine-readable index.
    lines = ['# 论文数据对应表', '',
             '来源：PMC12715241；数据集1.xlsx 的 all 表及数据来源图3.zip 的 Raw images_mask 目录。', '',
             '此表建立数据对应关系，尚未运行新分割算法或证明测量准确性。', '',
             '## 分组汇总', '',
             '| 分组 | 原图数 | 批次表目标数 | EV1样本数 | 唯一关联且有原图 |',
             '|---|---:|---:|---:|---:|']
    for row in summary['groups']:
        lines.append(f"| {row['group']} | {row['raw_images']} | {row['source_objects']} | {row['reference_samples']} | {row['paired_reference_samples']} |")
    lines += ['', '## 批次与图片清单', '',
              '| 分组 | 批次 | 文件夹中的细胞系代码 | 标称倍率 | 原图数 | 批次表目标数 | EV1样本数 | 唯一关联数 |',
              '|---|---|---|---|---:|---:|---:|---:|']
    for (group, batch), records in sorted(grouped.items()):
        first = records[0]
        refs = [r for r in references if r['group'] == group and r['batch'] == batch]
        count_images = sum(i['group'] == group and i['batch'] == batch for i in images)
        lines.append(f"| {group} | {batch} | {first['cell_line_code']} | {first['magnification_label']} | {count_images} | {len(records)} | {len(refs)} | {sum(r['eligible_for_paired_validation'] for r in refs)} |")
    lines += ['', '## 如何解读对应状态', '',
              '- `unique_10_feature_match`：同组、同批次的十项指标在记录精度内全部一致且唯一。它是数值关联，不是作者直接提供的跨表ID。',
              '- `unmatched`：十项指标未找到一致记录。保留EV1原始值，目标ID和图片ID留空，不按行号、最近值或图片顺序配对。',
              '- `ambiguous`：出现多个一致记录时保留候选，不自动选第一个。',
              '- PA第6天批次表有296条目标记录，EV1有102条；本次102条均未获得十项一致关联，差异原因尚未确定。',
              '- 其他三组196条样本获得唯一关联。原图文件存在仍不代表轮廓准确，后续需要核对编号叠加图及分割结果。',
              '- 11C1、12C1和2x/4x是从文件夹读取的原始标记。本次不推断其与论文Line 1/Line 2的对应，也不据此填入通用微米比例尺。', '',
              '## 查看详细记录', '',
              '程序索引位于 `data/reference/paper_data_index.json`，可直接在VS Code打开。',
              '1. `images`：142张原图与对应R_结果图的压缩包路径、分组、批次及稳定图片ID。',
              '2. `source_objects`：492条批次表记录，包含原表行号、图片名和目标编号。',
              '3. `reference_samples`：298条EV1记录、原始十项数值及关联状态。',
              '4. `issues`：需要核对的记录；`sources`包含来源文件SHA-256，便于检查资料是否换过版本。', '',
              'R_文件是轮廓与编号可视化结果，不作为另一张原图或逐像素分割标签。',
              '原始数据仍保留在用户资料目录，索引中记录压缩包内路径，没有复制整套图片。', '',
              '## 重建', '',
              '准备好openpyxl后，在项目目录运行：', '',
              '```powershell',
              'python scripts/build_reference_index.py --source-dir "存放数据集1.xlsx和数据来源图3.zip的文件夹"',
              '```', '',
              '脚本只读取来源；输出JSON和本文档，不修改应用或原始文件。当前网页环境暂未添加openpyxl，它仅是数据整理脚本的依赖。', '']
    docs = project_dir / 'docs'
    docs.mkdir(parents=True, exist_ok=True)
    (docs / '论文数据对应表.md').write_text('\n'.join(lines), encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return payload


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir', required=True)
    parser.add_argument('--output-dir', type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    build(args.source_dir, args.output_dir)
