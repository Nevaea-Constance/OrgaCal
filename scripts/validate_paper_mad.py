"""Read EV3/EV4 and compare explicit MAD conventions against author EV1 data."""
import argparse
import json
from pathlib import Path
import sys
import numpy as np
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from consistency import relative_mad

MAPPING={'area[um2]':'area','perimeter[um]':'perimeter','avg_radius[um]':'paper_radius',
         'roundness[%]':'paper_roundness_pct','maxFeret[um]':'max_feret','minFeret[um]':'min_feret',
         'Mean curvature[um-1]':'mean_curvature','Std curvature[um-1]':'std_curvature',
         'Std curvature x R0':'std_curvature_times_radius','DNE':'paper_dne'}
MAPPING.update({'area':'area','perimeter':'perimeter','avg_radius':'paper_radius','roundness':'paper_roundness_pct',
                'maxferet':'max_feret','minferet':'min_feret','mean_curvature':'mean_curvature',
                'std_curvature':'std_curvature','std_curvaturexr0':'std_curvature_times_radius'})

def main(source):
    index=json.loads((ROOT/'data/reference/paper_data_index.json').read_text(encoding='utf-8'))
    inverse={'area':'area_um2','perimeter':'perimeter_um','paper_radius':'paper_radius_um',
             'paper_roundness_pct':'paper_roundness_pct','max_feret':'max_feret_um','min_feret':'min_feret_um',
             'mean_curvature':'mean_curvature_per_um','std_curvature':'std_curvature_per_um',
             'std_curvature_times_radius':'std_curvature_times_radius','paper_dne':'paper_dne'}
    records=[{'condition':r['group'].split('_')[0],'batch':r['batch'],'day':str(r['day']),
              **{k:r['metrics'][v] for k,v in inverse.items()}} for r in index['reference_samples']]
    rows=[]
    for filename,day,sheet in [('数据集3.xlsx',6,'MAD_Day6'),('数据集4.xlsx',90,'MAD_Day90')]:
        workbook=load_workbook(Path(source)/filename,read_only=True,data_only=True)
        source_rows=list(workbook[sheet].iter_rows(min_row=2,values_only=True))
        print('SOURCE_VARIABLES',sorted({row[0] for row in source_rows}))
        for source_row,(variable,batch,group,value) in enumerate(source_rows,2):
            if variable not in MAPPING:
                raise ValueError(f'Unknown variable {variable}')
            feature=MAPPING[variable]
            x=[r[feature] for r in records if r['condition']==group and r['batch']==batch and r['day']==str(day)]
            r=relative_mad(x)
            variants={'raw_abs_mad':r['mad'],'r_scaled_abs_mad':r['mad']*1.4826,
                      'relative_mad_ratio':None if r['relative_mad_pct'] is None else r['relative_mad_pct']/100,
                      'r_scaled_relative_ratio':None if r['relative_mad_pct'] is None else r['relative_mad_pct']/100*1.4826}
            matches=[name for name,n in variants.items() if n is not None and np.isclose(n,value,rtol=1e-8,atol=1e-10)]
            rows.append({'file':filename,'sheet':sheet,'source_row':source_row,'day':day,'batch':batch,'condition':group,
                         'variable':variable,'reference':value,'n':r['n'],'computed':variants,'matched_conventions':matches})
    counts={key:sum(key in r['matched_conventions'] for r in rows) for key in rows[0]['computed']}
    output={'rows_checked':len(rows),'match_counts':counts,'all_rows_matched':all(r['matched_conventions'] for r in rows),'rows':rows}
    path=ROOT/'data/reference/mad_validation.json'
    path.write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in output.items() if k!='rows'},ensure_ascii=False,indent=2))
    for r in rows:
        if not r['matched_conventions']:print('UNMATCHED',r['variable'],r['condition'],r['batch'],r['day'],r['reference'],r['computed'])

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--source-dir',required=True)
    main(parser.parse_args().source_dir)
